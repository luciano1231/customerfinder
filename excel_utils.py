import os
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'busquedas_excel')

ESTADOS = {"nuevo": "Nuevo", "en_cola": "En cola", "enviado": "Enviado", "sin_whatsapp": "Sin WhatsApp",
           "error": "Error", "respondio": "Respondió", "no_interesa": "No le interesa"}


def sanitize_filename(name):
    """Remueve caracteres no válidos para nombres de archivo en Windows."""
    return re.sub(r'[\\/*?:"<>|]', "", name).strip()


def create_excel_file(title, rows, output_dir=EXCEL_DIR):
    """Excel con los locales (ordenados como vienen) y links directos a WhatsApp, Instagram y Maps."""
    os.makedirs(output_dir, exist_ok=True)
    filename = f"{sanitize_filename(title)} - {datetime.now().strftime('%Y-%m-%d %H%M')}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Locales"
    headers = ["Nombre", "Categoría", "Reseñas", "Rating", "Teléfono", "Tipo tel.", "WhatsApp", "Web",
               "Instagram / link", "Nocturno", "Dirección", "Horario", "Google Maps", "Estado", "Notas"]
    ws.append(headers)
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    for c in ws[1]:
        c.fill = header_fill
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    link = Font(color="2563EB", underline="single")
    wa = Font(color="16A34A", underline="single", bold=True)
    border = Border(*(Side(style='thin', color='E2E8F0'),) * 4)

    for r, p in enumerate(rows, start=2):
        social = p.get("instagram") or (p.get("web") if p.get("web_tipo") != "Web propia" else "")
        values = [p.get("name"), p.get("categoria"), p.get("reviews") or 0, p.get("rating"), p.get("phone"),
                  p.get("tel_tipo"), p.get("wa_link") or "Sin WhatsApp", p.get("web_tipo"), social,
                  "Sí" if p.get("nocturno") else "No", p.get("address"), " | ".join(p.get("horarios") or []),
                  "Ver en Maps" if p.get("maps_url") else "", ESTADOS.get(p.get("estado"), p.get("estado")),
                  p.get("notas")]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
        if p.get("wa_link"):
            ws.cell(row=r, column=7).hyperlink = p["wa_link"]
            ws.cell(row=r, column=7).font = wa
        if social:
            ws.cell(row=r, column=9).hyperlink = social
            ws.cell(row=r, column=9).font = link
        if p.get("maps_url"):
            ws.cell(row=r, column=13).hyperlink = p["maps_url"]
            ws.cell(row=r, column=13).font = link

    widths = [34, 22, 9, 7, 17, 10, 30, 18, 36, 9, 38, 50, 13, 14, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return filepath, filename
